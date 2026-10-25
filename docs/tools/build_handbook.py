"""Build the handbook Word edition and legacy documentation entry points.

This optional documentation tool requires python-docx. It does not run the app,
connect to a database, change startup settings or refresh a data snapshot.
"""
from pathlib import Path
import argparse
import copy
import re
from urllib.parse import quote
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

DOCS = Path(__file__).resolve().parents[1]
SOURCE = DOCS / 'UP-SELF-Developer-Handbook.md'
OUTPUT = DOCS / 'UP-SELF-Developer-Handbook.docx'


def field(paragraph, instruction, fallback='1'):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    code = OxmlElement('w:instrText'); code.set(qn('xml:space'), 'preserve'); code.text = ' '+instruction+' '
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t'); text.text = fallback
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for item in [begin, code, separate, text, end]: run._r.append(item)


def bookmark(paragraph, identifier, name):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(identifier));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(identifier))
    paragraph._p.insert(0,start);paragraph._p.append(end)


def add_link(paragraph, label, target=None, anchor=None):
    h=OxmlElement('w:hyperlink')
    if anchor: h.set(qn('w:anchor'),anchor)
    else:
        h.set(qn('r:id'),paragraph.part.relate_to(target,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'1B4C70');props.append(color)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'20');props.append(size)
    run.append(props);text=OxmlElement('w:t');text.text=label.strip('`');run.append(text);h.append(run);paragraph._p.append(h)


def inline(paragraph, text, table=False):
    pattern=r'(\[[^\]]+\]\([^\)]+\)|`[^`]+`|\*\*[^*]+\*\*)'
    for token in re.split(pattern,text):
        if not token:continue
        link=re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)',token)
        if link:
            add_link(paragraph,link[1],link[2]);continue
        code=token.startswith('`') and token.endswith('`')
        bold=token.startswith('**') and token.endswith('**')
        run=paragraph.add_run(token[1:-1] if code else token[2:-2] if bold else token)
        if code:run.font.name='Consolas';run.font.size=Pt(8.1 if table else 9)
        if bold:run.bold=True


def cell_margins(cell):
    tcPr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
    for edge,value in [('top',75),('bottom',75),('left',95),('right',95)]:
        node=OxmlElement('w:'+edge);node.set(qn('w:w'),str(value));node.set(qn('w:type'),'dxa');m.append(node)
    tcPr.append(m)


def add_table(doc, rows):
    table=doc.add_table(rows=1,cols=len(rows[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    cols=len(rows[0]);headers=rows[0]
    if headers[0]=='FR': weights=[.09,.24,.14,.53]
    elif headers[:2]==['Method','Path after prefix']:weights=[.09,.40,.37,.14]
    elif headers[0]=='Resource or action':weights=[.28,.17,.25,.14,.16]
    elif headers[0]=='Option':weights=[.085,.225,.225,.225,.24]
    elif cols==2:weights=[.32,.68]
    elif cols==3:weights=[.30,.24,.46]
    elif cols==4:weights=[.25,.23,.23,.29]
    elif cols==5:weights=[.26,.31,.12,.17,.14]
    else:weights=[1/cols]*cols
    available=6.9
    for col,weight in zip(table.columns,weights):col.width=Inches(available*weight)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    table._tbl.tblPr.append(borders)
    for index,row in enumerate(rows):
        cells=table.rows[0].cells if index==0 else table.add_row().cells
        for cno,value in enumerate(row):
            cell=cells[cno];cell.width=Inches(available*weights[cno]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell_margins(cell)
            p=cell.paragraphs[0];p.style=doc.styles['Table Text'];inline(p,value,True)
            p.paragraph_format.keep_with_next=index==0
            if index==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E8EDF2');cell._tc.get_or_add_tcPr().append(fill)
                for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(0,0,0)
            elif index%2==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'F7F8FA');cell._tc.get_or_add_tcPr().append(fill)
            if len(value)<18 and (cno==0 and headers[0] in ('FR','Method') or value in ('Completed','Partial','Not implemented','Data','Table','File')):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        trPr=table.rows[index]._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
        if index==0:
            repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trPr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)


def build():
    text=SOURCE.read_text(encoding='utf-8');lines=text.splitlines();doc=Document()
    section=doc.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.left_margin=Inches(.8);section.right_margin=Inches(.8);section.top_margin=Inches(.8);section.bottom_margin=Inches(.75)
    section.header_distance=Inches(.3);section.footer_distance=Inches(.3)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.12
    for name,size in [('Title',27),('Subtitle',12),('Heading 1',17),('Heading 2',12),('Heading 3',11)]:
        style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_before=Pt(14);style.paragraph_format.space_after=Pt(7)
        style.paragraph_format.keep_with_next=True
    doc.styles['Heading 1'].paragraph_format.page_break_before=True
    doc.styles['Title'].paragraph_format.space_before=Pt(40)
    for name in ['Code Text','Table Text','TOC Entry']:
        if name not in doc.styles:doc.styles.add_style(name,1)
    code=doc.styles['Code Text'];code.font.name='Consolas';code.font.size=Pt(8.2);code.paragraph_format.space_before=Pt(3);code.paragraph_format.space_after=Pt(7);code.paragraph_format.line_spacing=1.06
    code.paragraph_format.widow_control=False
    table_style=doc.styles['Table Text'];table_style.font.name='Arial';table_style.font.size=Pt(8.7);table_style.paragraph_format.space_after=Pt(2);table_style.paragraph_format.line_spacing=1.08
    for style in doc.styles:
        if style.type==1:
            style.element.get_or_add_rPr()
            for border in list(style.element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Page ');field(footer,'PAGE');footer.add_run(' of ');field(footer,'NUMPAGES')
    for r in footer.runs:r.font.name='Arial';r.font.size=Pt(8)
    # A real TOC field is updated during final rendering and remains refreshable in Word.
    title=lines[0].removeprefix('# ');doc.add_paragraph(title,'Title')
    body_start=next(i for i,line in enumerate(lines) if line.startswith('## '))
    paragraphs='\n'.join(lines[1:body_start]).strip().split('\n\n')
    for i,paragraph in enumerate(paragraphs):
        p=doc.add_paragraph(style='Subtitle' if i<2 else 'Normal');inline(p,paragraph)
    doc.add_page_break()
    p=doc.add_paragraph('Contents','Title');p.paragraph_format.space_before=Pt(0)
    p=doc.add_paragraph();field(p,'TOC \\o "1-2" \\h \\z \\u','')
    # Publication tooling may use this invisible insertion bookmark to populate
    # the native index. Word users can update the TOC field directly.
    bookmark(doc.add_paragraph(),10000,'handbook_toc')
    settings=doc.settings.element
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    bid=1;i=body_start
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line.startswith('```'):
            if doc.paragraphs and doc.paragraphs[-1].style.name=='Normal':
                doc.paragraphs[-1].paragraph_format.keep_with_next=True
            code_lines=[];i+=1
            while i<len(lines) and not lines[i].startswith('```'):code_lines.append(lines[i]);i+=1
            p=doc.add_paragraph('\n'.join(code_lines),'Code Text')
            p.paragraph_format.keep_together=len(code_lines)<=14
            i+=1;continue
        if line.startswith('|'):
            if doc.paragraphs and doc.paragraphs[-1].style.name=='Normal':
                doc.paragraphs[-1].paragraph_format.keep_with_next=True
            table_lines=[]
            while i<len(lines) and lines[i].startswith('|'):
                if not re.fullmatch(r'[|:\-\s]+',lines[i]):table_lines.append([v.strip() for v in lines[i].strip().strip('|').split('|')])
                i+=1
            if table_lines:add_table(doc,table_lines)
            continue
        heading=re.match(r'^(#{2,4}) (.*)',line)
        if heading:
            level=len(heading[1])-1;p=doc.add_paragraph(heading[2],'Heading '+str(level));bookmark(p,bid,'section_'+str(bid));bid+=1;i+=1;continue
        bullet=re.match(r'^- (.*)',line);number=re.match(r'^\d+\. (.*)',line)
        if bullet or number:
            p=doc.add_paragraph(style='List Bullet' if bullet else 'List Number');inline(p,(bullet or number)[1]);i+=1;continue
        parts=[line];i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||```|- |\d+\. )',lines[i]):parts.append(lines[i]);i+=1
        p=doc.add_paragraph();inline(p,' '.join(parts))
    doc.core_properties.title='UP SELF Developer Handbook';doc.core_properties.subject='UP-SELF development and operations'
    doc.core_properties.author='UP-SELF';doc.core_properties.last_modified_by='UP-SELF';doc.core_properties.comments=''
    doc.core_properties.keywords='Java Vue development permissions migrations deployment'
    OUTPUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUTPUT)
    print('Created',OUTPUT)


def compatibility():
    text=SOURCE.read_text(encoding='utf-8')
    destinations={
      '跨电脑环境安装与启动教程.md':('Local setup and data migration','#5-local-setup-and-data-migration'),
      'Mac环境安装与迁移启动教程.md':('Mac setup and startup','#58-m1-mac-prerequisites')}
    for name,(title,anchor) in destinations.items():
        (DOCS/name).write_text('# '+title+'\n\nThis compatibility entry preserves an existing script or metadata path. Read the [Developer Handbook](UP-SELF-Developer-Handbook.md'+anchor+') for the maintained instructions. Do not edit a second set of setup steps here.\n',encoding='utf-8')
    section=text.split('### A2 Complete requirement matrix\n',1)[1].split('\n### A3',1)[0].strip()
    (DOCS/'当前61个FR完成状态与报告差异.md').write_text('# Functional requirement compatibility index\n\nGenerated from Appendix A of the [Developer Handbook](UP-SELF-Developer-Handbook.md#appendix-a-functional-requirements). This file preserves the path read by the existing Linux preparation verifier. Edit the handbook, then regenerate this index.\n\n'+section+'\n',encoding='utf-8')


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--compatibility-only',action='store_true');args=parser.parse_args()
    if not args.compatibility_only:build()
    compatibility()
